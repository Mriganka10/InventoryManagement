from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from io import StringIO
import csv
import json
import os
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from reportlab.graphics.barcode import code128
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError

from .models import (
    Appointment, AuditLog, Branch, Customer, InventoryItem, Invoice, InvoiceLine,
    JobCard, JobCardItem, Payment, PurchaseOrder, PurchaseOrderLine, SerialNumber,
    ServicePackage, ServicePackageItem, SessionLocal, StockBalance, StockBatch,
    StockCount, StockCountLine, StockMovement, Supplier, User, Vehicle, Workshop,
    WorkshopUser, utcnow,
)

ZERO = Decimal("0")


def as_dict(row: Any) -> dict:
    result = {}
    for column in row.__table__.columns:
        value = getattr(row, column.name)
        if isinstance(value, Decimal):
            value = float(value)
        elif isinstance(value, (date, datetime)):
            value = value.isoformat()
        result[column.name] = value
    return result


def context_for(email: str) -> dict:
    with SessionLocal() as db:
        result = db.execute(
            select(User, WorkshopUser, Workshop)
            .join(WorkshopUser, WorkshopUser.user_id == User.id)
            .join(Workshop, Workshop.id == WorkshopUser.workshop_id)
            .where(User.email == email, WorkshopUser.active.is_(True))
        ).first()
        if not result:
            raise HTTPException(403, "No active workshop is linked to this account.")
        user, membership, workshop = result
        return {"user": as_dict(user), "membership": as_dict(membership), "workshop": as_dict(workshop)}


def _number(value: Any, default: str = "0") -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except Exception as exc:
        raise HTTPException(400, f"Invalid numeric value: {value}") from exc


def _audit(db, ctx: dict, action: str, entity: Any, before: dict | None = None) -> None:
    db.add(AuditLog(
        workshop_id=ctx["workshop"]["id"], user_id=ctx["user"]["id"], action=action,
        entity_type=entity.__tablename__, entity_id=getattr(entity, "id", None),
        before_data=before or {}, after_data=as_dict(entity),
    ))


def bootstrap(ctx: dict) -> dict:
    wid = ctx["workshop"]["id"]
    with SessionLocal() as db:
        branches = db.scalars(select(Branch).where(Branch.workshop_id == wid, Branch.active.is_(True)).order_by(Branch.name)).all()
        return {**ctx, "branches": [as_dict(x) for x in branches]}


ENTITY_MODELS = {
    "branches": Branch,
    "suppliers": Supplier,
    "customers": Customer,
    "vehicles": Vehicle,
    "items": InventoryItem,
    "service-packages": ServicePackage,
}


def list_entities(ctx: dict, kind: str, search: str = "", limit: int = 250) -> list[dict]:
    model = ENTITY_MODELS.get(kind)
    if not model:
        raise HTTPException(404, "Unknown entity type.")
    wid = ctx["workshop"]["id"]
    with SessionLocal() as db:
        query = select(model).where(model.workshop_id == wid)
        if search:
            terms = []
            for attr in ("name", "sku", "barcode", "phone", "registration_number", "code"):
                if hasattr(model, attr):
                    terms.append(getattr(model, attr).ilike(f"%{search}%"))
            if terms:
                query = query.where(or_(*terms))
        rows = db.scalars(query.order_by(model.id.desc()).limit(max(1, min(limit, 1000)))).all()
        result = [as_dict(row) for row in rows]
        if kind == "items":
            balance_rows = db.execute(
                select(StockBalance.item_id, func.sum(StockBalance.quantity_on_hand), func.sum(StockBalance.quantity_reserved))
                .where(StockBalance.workshop_id == wid).group_by(StockBalance.item_id)
            ).all()
            balances = {item_id: (float(on_hand or 0), float(reserved or 0)) for item_id, on_hand, reserved in balance_rows}
            for item in result:
                on_hand, reserved = balances.get(item["id"], (0.0, 0.0))
                item.update(quantity_on_hand=on_hand, quantity_reserved=reserved, quantity_available=on_hand - reserved)
        return result


def list_operations(ctx: dict, kind: str, limit: int = 200) -> list[dict]:
    models = {"appointments": Appointment, "jobs": JobCard, "invoices": Invoice, "purchases": PurchaseOrder, "payments": Payment}
    model = models.get(kind)
    if not model:
        raise HTTPException(404, "Unknown operation type.")
    with SessionLocal() as db:
        rows = db.scalars(select(model).where(model.workshop_id == ctx["workshop"]["id"]).order_by(model.id.desc()).limit(min(limit, 500))).all()
        return [as_dict(row) for row in rows]


def create_entity(ctx: dict, kind: str, payload: dict) -> dict:
    model = ENTITY_MODELS.get(kind)
    if not model:
        raise HTTPException(404, "Unknown entity type.")
    wid = ctx["workshop"]["id"]
    allowed = {column.name for column in model.__table__.columns} - {"id", "workshop_id", "created_at", "updated_at"}
    values = {key: value for key, value in payload.items() if key in allowed and value != ""}
    numeric = {"purchase_price", "selling_price", "tax_rate", "min_stock", "reorder_quantity", "lead_time_days", "rating", "model_year", "odometer_km", "labour_price"}
    for key in numeric & values.keys():
        values[key] = _number(values[key])
    if kind == "branches":
        values.setdefault("code", f"BR-{int(utcnow().timestamp())}")
    if kind == "suppliers":
        values.setdefault("code", f"SUP-{int(utcnow().timestamp())}")
    if kind == "items":
        values.setdefault("sku", f"ITEM-{int(utcnow().timestamp())}")
        values.setdefault("barcode", values["sku"])
    try:
        with SessionLocal.begin() as db:
            if kind == "vehicles":
                customer = db.scalar(select(Customer).where(Customer.id == int(values.get("customer_id") or 0), Customer.workshop_id == wid))
                if not customer:
                    raise HTTPException(404, "Customer not found in this workshop.")
            if kind == "items" and values.get("preferred_supplier_id"):
                supplier = db.scalar(select(Supplier).where(Supplier.id == int(values["preferred_supplier_id"]), Supplier.workshop_id == wid))
                if not supplier:
                    raise HTTPException(404, "Supplier not found in this workshop.")
            entity = model(workshop_id=wid, **values)
            db.add(entity); db.flush(); _audit(db, ctx, "create", entity); result = as_dict(entity)
        return result
    except IntegrityError as exc:
        raise HTTPException(409, "A record with the same unique code already exists.") from exc


def create_customer_vehicle(ctx: dict, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    customer_data = payload.get("customer") or {}
    vehicle_data = payload.get("vehicle") or {}
    if not customer_data.get("name") or not vehicle_data.get("registration_number"):
        raise HTTPException(400, "Customer name and vehicle registration number are required.")
    with SessionLocal.begin() as db:
        customer = Customer(workshop_id=wid, **{k: v for k, v in customer_data.items() if k in {"name", "email", "phone", "address", "tax_id", "notes"}})
        db.add(customer); db.flush()
        vehicle = Vehicle(workshop_id=wid, customer_id=customer.id, **{k: v for k, v in vehicle_data.items() if k in {"registration_number", "vin", "make", "model", "model_year", "fuel_type", "engine_number", "odometer_km"}})
        db.add(vehicle); db.flush(); _audit(db, ctx, "create", customer); _audit(db, ctx, "create", vehicle)
        return {"customer": as_dict(customer), "vehicle": as_dict(vehicle)}


def _balance(db, wid: int, branch_id: int, item_id: int, lock: bool = True) -> StockBalance:
    query = select(StockBalance).where(
        StockBalance.workshop_id == wid, StockBalance.branch_id == branch_id, StockBalance.item_id == item_id
    )
    if lock and not str(db.bind.url).startswith("sqlite"):
        query = query.with_for_update()
    balance = db.scalar(query)
    if not balance:
        balance = StockBalance(workshop_id=wid, branch_id=branch_id, item_id=item_id)
        db.add(balance); db.flush()
    return balance


def stock_movement(ctx: dict, payload: dict) -> dict:
    wid, uid = ctx["workshop"]["id"], ctx["user"]["id"]
    branch_id, item_id = int(payload.get("branch_id") or 0), int(payload.get("item_id") or 0)
    movement_type = str(payload.get("movement_type") or "adjustment")
    quantity = _number(payload.get("quantity"))
    if not branch_id or not item_id or quantity == 0:
        raise HTTPException(400, "Branch, item, and a non-zero quantity are required.")
    positive_types = {"receipt", "return", "transfer_in", "opening", "adjustment_in"}
    negative_types = {"issue", "damage", "expiry", "transfer_out", "adjustment_out"}
    if movement_type not in positive_types | negative_types:
        raise HTTPException(400, "Unsupported stock movement type.")
    delta = abs(quantity) if movement_type in positive_types else -abs(quantity)
    with SessionLocal.begin() as db:
        item = db.scalar(select(InventoryItem).where(InventoryItem.id == item_id, InventoryItem.workshop_id == wid))
        branch = db.scalar(select(Branch).where(Branch.id == branch_id, Branch.workshop_id == wid))
        if not item or not branch:
            raise HTTPException(404, "Item or branch not found.")
        balance = _balance(db, wid, branch_id, item_id)
        if balance.quantity_on_hand + delta < 0:
            raise HTTPException(409, "This transaction would make stock negative.")
        balance.quantity_on_hand += delta
        for key in ("rack", "shelf", "bin_code"):
            if payload.get(key) is not None:
                setattr(balance, key, str(payload[key]))
        batch = None
        if payload.get("batch_number"):
            batch = db.scalar(select(StockBatch).where(
                StockBatch.workshop_id == wid, StockBatch.branch_id == branch_id,
                StockBatch.item_id == item_id, StockBatch.batch_number == str(payload["batch_number"]),
            ))
            if not batch:
                batch = StockBatch(workshop_id=wid, branch_id=branch_id, item_id=item_id, batch_number=str(payload["batch_number"]))
                db.add(batch); db.flush()
            if payload.get("expires_on"):
                batch.expires_on = date.fromisoformat(str(payload["expires_on"]))
            if payload.get("warranty_until"):
                batch.warranty_until = date.fromisoformat(str(payload["warranty_until"]))
            if batch.quantity + delta < 0:
                raise HTTPException(409, "This transaction would make batch stock negative.")
            batch.quantity += delta
        serials = payload.get("serial_numbers") or []
        if item.track_serial and abs(delta) != len(serials):
            raise HTTPException(400, "Provide one serial number for each serialized unit.")
        for serial in serials:
            existing = db.scalar(select(SerialNumber).where(SerialNumber.workshop_id == wid, SerialNumber.serial_number == str(serial)))
            if delta > 0 and existing:
                raise HTTPException(409, f"Serial number already exists: {serial}")
            if delta > 0:
                db.add(SerialNumber(workshop_id=wid, branch_id=branch_id, item_id=item_id, batch_id=batch.id if batch else None, serial_number=str(serial)))
            elif existing:
                existing.status = "issued"
        movement = StockMovement(
            workshop_id=wid, branch_id=branch_id, item_id=item_id, batch_id=batch.id if batch else None,
            movement_type=movement_type, quantity=delta, unit_cost=_number(payload.get("unit_cost", item.purchase_price)),
            reference_type=str(payload.get("reference_type") or "manual"), reference_id=payload.get("reference_id"),
            reason=str(payload.get("reason") or ""), performed_by=uid,
        )
        db.add(movement); db.flush(); _audit(db, ctx, "stock_movement", movement)
        return {"movement": as_dict(movement), "balance": as_dict(balance)}


def transfer_stock(ctx: dict, payload: dict) -> dict:
    source, target = int(payload.get("from_branch_id") or 0), int(payload.get("to_branch_id") or 0)
    if not source or not target or source == target:
        raise HTTPException(400, "Choose two different branches.")
    wid, uid = ctx["workshop"]["id"], ctx["user"]["id"]
    item_id, quantity = int(payload.get("item_id") or 0), abs(_number(payload.get("quantity")))
    if not item_id or quantity <= 0:
        raise HTTPException(400, "Item and positive quantity are required.")
    with SessionLocal.begin() as db:
        branches = db.scalars(select(Branch).where(Branch.workshop_id == wid, Branch.id.in_([source, target]))).all()
        item = db.scalar(select(InventoryItem).where(InventoryItem.id == item_id, InventoryItem.workshop_id == wid))
        if len(branches) != 2 or not item:
            raise HTTPException(404, "Item or branch not found in this workshop.")
        from_balance, to_balance = _balance(db, wid, source, item_id), _balance(db, wid, target, item_id)
        if from_balance.quantity_on_hand - from_balance.quantity_reserved < quantity:
            raise HTTPException(409, "Insufficient available stock for transfer.")
        from_balance.quantity_on_hand -= quantity; to_balance.quantity_on_hand += quantity
        reason = str(payload.get("reason") or "Branch transfer")
        outgoing = StockMovement(workshop_id=wid, branch_id=source, item_id=item_id, movement_type="transfer_out", quantity=-quantity, unit_cost=item.purchase_price, reference_type="branch_transfer", reason=reason, performed_by=uid)
        incoming = StockMovement(workshop_id=wid, branch_id=target, item_id=item_id, movement_type="transfer_in", quantity=quantity, unit_cost=item.purchase_price, reference_type="branch_transfer", reason=reason, performed_by=uid)
        db.add_all([outgoing, incoming]); db.flush(); outgoing.reference_id = incoming.id; incoming.reference_id = outgoing.id
        _audit(db, ctx, "transfer_out", outgoing); _audit(db, ctx, "transfer_in", incoming)
        return {"outgoing": as_dict(outgoing), "incoming": as_dict(incoming)}


def create_appointment(ctx: dict, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    required = ["branch_id", "customer_id", "vehicle_id", "scheduled_at"]
    if any(not payload.get(key) for key in required):
        raise HTTPException(400, "Branch, customer, vehicle, and scheduled time are required.")
    scheduled = datetime.fromisoformat(str(payload["scheduled_at"]).replace("Z", "+00:00"))
    with SessionLocal.begin() as db:
        branch = db.scalar(select(Branch).where(Branch.id == int(payload["branch_id"]), Branch.workshop_id == wid))
        customer = db.scalar(select(Customer).where(Customer.id == int(payload["customer_id"]), Customer.workshop_id == wid))
        vehicle = db.scalar(select(Vehicle).where(Vehicle.id == int(payload["vehicle_id"]), Vehicle.workshop_id == wid, Vehicle.customer_id == int(payload["customer_id"])))
        package_id = int(payload.get("service_package_id") or 0)
        package = db.scalar(select(ServicePackage).where(ServicePackage.id == package_id, ServicePackage.workshop_id == wid)) if package_id else None
        if not branch or not customer or not vehicle or (package_id and not package):
            raise HTTPException(404, "A selected branch, customer, vehicle, or service package is not valid for this workshop.")
        appointment = Appointment(
            workshop_id=wid, branch_id=int(payload["branch_id"]), customer_id=int(payload["customer_id"]),
            vehicle_id=int(payload["vehicle_id"]), service_package_id=payload.get("service_package_id") or None,
            scheduled_at=scheduled, complaint=str(payload.get("complaint") or ""),
        )
        db.add(appointment); db.flush()
        job = JobCard(
            workshop_id=wid, branch_id=appointment.branch_id, appointment_id=appointment.id,
            customer_id=appointment.customer_id, vehicle_id=appointment.vehicle_id,
            job_number=f"JOB-{utcnow():%Y%m%d}-{appointment.id:05d}", status="scheduled", complaint=appointment.complaint,
        )
        db.add(job); db.flush()
        shortage = []
        if appointment.service_package_id:
            recipe = db.scalars(select(ServicePackageItem).where(
                ServicePackageItem.workshop_id == wid, ServicePackageItem.service_package_id == appointment.service_package_id
            )).all()
            for part in recipe:
                balance = _balance(db, wid, appointment.branch_id, part.item_id)
                available = balance.quantity_on_hand - balance.quantity_reserved
                reserve = min(max(available, ZERO), part.quantity)
                balance.quantity_reserved += reserve
                item = db.get(InventoryItem, part.item_id)
                db.add(JobCardItem(workshop_id=wid, job_card_id=job.id, item_id=part.item_id, quantity_reserved=reserve, unit_price=item.selling_price if item else 0))
                if reserve < part.quantity:
                    shortage.append({"item_id": part.item_id, "required": float(part.quantity), "reserved": float(reserve)})
        _audit(db, ctx, "create_with_auto_reservation", appointment)
        return {"appointment": as_dict(appointment), "job_card": as_dict(job), "shortages": shortage}


def set_service_package_items(ctx: dict, package_id: int, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    lines = payload.get("items") or []
    with SessionLocal.begin() as db:
        package = db.scalar(select(ServicePackage).where(ServicePackage.id == package_id, ServicePackage.workshop_id == wid))
        if not package:
            raise HTTPException(404, "Service package not found.")
        db.query(ServicePackageItem).filter(ServicePackageItem.workshop_id == wid, ServicePackageItem.service_package_id == package_id).delete()
        for line in lines:
            item = db.scalar(select(InventoryItem).where(InventoryItem.id == int(line["item_id"]), InventoryItem.workshop_id == wid))
            if not item:
                raise HTTPException(404, f"Inventory item not found: {line.get('item_id')}")
            db.add(ServicePackageItem(workshop_id=wid, service_package_id=package_id, item_id=item.id, quantity=_number(line.get("quantity", 1))))
        _audit(db, ctx, "set_recipe", package)
        return {"service_package": as_dict(package), "items": lines}


def create_purchase(ctx: dict, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    lines = payload.get("lines") or []
    if not lines:
        raise HTTPException(400, "At least one purchase line is required.")
    with SessionLocal.begin() as db:
        branch = db.scalar(select(Branch).where(Branch.id == int(payload.get("branch_id") or 0), Branch.workshop_id == wid))
        supplier = db.scalar(select(Supplier).where(Supplier.id == int(payload.get("supplier_id") or 0), Supplier.workshop_id == wid))
        if not supplier or not branch:
            raise HTTPException(404, "Supplier or branch not found in this workshop.")
        po = PurchaseOrder(
            workshop_id=wid, branch_id=int(payload["branch_id"]), supplier_id=supplier.id,
            po_number=str(payload.get("po_number") or f"PO-{utcnow():%Y%m%d%H%M%S}"), status="ordered",
            ordered_on=date.today(), expected_on=date.fromisoformat(payload["expected_on"]) if payload.get("expected_on") else date.today() + timedelta(days=supplier.lead_time_days),
            notes=str(payload.get("notes") or ""),
        )
        db.add(po); db.flush(); subtotal = ZERO; tax_total = ZERO
        for raw in lines:
            item = db.scalar(select(InventoryItem).where(InventoryItem.id == int(raw["item_id"]), InventoryItem.workshop_id == wid))
            if not item:
                raise HTTPException(404, "Purchase item not found.")
            quantity, cost, tax = _number(raw.get("quantity")), _number(raw.get("unit_cost", item.purchase_price)), _number(raw.get("tax_rate", item.tax_rate))
            db.add(PurchaseOrderLine(workshop_id=wid, purchase_order_id=po.id, item_id=item.id, quantity_ordered=quantity, unit_cost=cost, tax_rate=tax))
            subtotal += quantity * cost; tax_total += quantity * cost * tax / 100
        po.subtotal, po.tax_total, po.total = subtotal, tax_total, subtotal + tax_total
        _audit(db, ctx, "create", po)
        return as_dict(po)


def receive_purchase(ctx: dict, purchase_id: int) -> dict:
    wid, uid = ctx["workshop"]["id"], ctx["user"]["id"]
    with SessionLocal.begin() as db:
        po = db.scalar(select(PurchaseOrder).where(PurchaseOrder.id == purchase_id, PurchaseOrder.workshop_id == wid).with_for_update())
        if not po:
            raise HTTPException(404, "Purchase order not found.")
        if po.status == "received":
            raise HTTPException(409, "Purchase order has already been received.")
        lines = db.scalars(select(PurchaseOrderLine).where(PurchaseOrderLine.purchase_order_id == po.id)).all()
        for line in lines:
            quantity = line.quantity_ordered - line.quantity_received
            if quantity <= 0:
                continue
            balance = _balance(db, wid, po.branch_id, line.item_id); balance.quantity_on_hand += quantity
            line.quantity_received += quantity
            db.add(StockMovement(workshop_id=wid, branch_id=po.branch_id, item_id=line.item_id, movement_type="receipt", quantity=quantity, unit_cost=line.unit_cost, reference_type="purchase_order", reference_id=po.id, reason=f"Received {po.po_number}", performed_by=uid))
        po.status, po.received_on = "received", date.today(); _audit(db, ctx, "receive", po)
        return as_dict(po)


def start_stock_count(ctx: dict, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    branch_id = int(payload.get("branch_id") or 0)
    counts = payload.get("counts") or []
    if not branch_id or not counts:
        raise HTTPException(400, "Branch and physical counts are required.")
    with SessionLocal.begin() as db:
        branch = db.scalar(select(Branch).where(Branch.id == branch_id, Branch.workshop_id == wid))
        if not branch:
            raise HTTPException(404, "Branch not found in this workshop.")
        stock_count = StockCount(workshop_id=wid, branch_id=branch_id, status="submitted", method="blind_count", counted_by=ctx["user"]["id"])
        db.add(stock_count); db.flush()
        for raw in counts:
            item_id, counted = int(raw["item_id"]), _number(raw.get("counted_quantity"))
            item = db.scalar(select(InventoryItem).where(InventoryItem.id == item_id, InventoryItem.workshop_id == wid))
            if not item:
                raise HTTPException(404, f"Inventory item not found: {item_id}")
            balance = _balance(db, wid, branch_id, item_id, lock=False)
            db.add(StockCountLine(workshop_id=wid, stock_count_id=stock_count.id, item_id=item_id, system_quantity=balance.quantity_on_hand, counted_quantity=counted, variance=counted - balance.quantity_on_hand, reason=str(raw.get("reason") or "")))
        _audit(db, ctx, "submit_blind_count", stock_count)
        return as_dict(stock_count)


def approve_stock_count(ctx: dict, count_id: int) -> dict:
    wid, uid = ctx["workshop"]["id"], ctx["user"]["id"]
    if ctx["membership"]["role"] not in {"owner", "manager"}:
        raise HTTPException(403, "Only an owner or manager can approve stock variances.")
    with SessionLocal.begin() as db:
        count = db.scalar(select(StockCount).where(StockCount.id == count_id, StockCount.workshop_id == wid).with_for_update())
        if not count or count.status != "submitted":
            raise HTTPException(409, "Submitted stock count not found.")
        lines = db.scalars(select(StockCountLine).where(StockCountLine.stock_count_id == count.id)).all()
        for line in lines:
            if line.variance == 0:
                continue
            balance = _balance(db, wid, count.branch_id, line.item_id); balance.quantity_on_hand += line.variance
            db.add(StockMovement(workshop_id=wid, branch_id=count.branch_id, item_id=line.item_id, movement_type="adjustment_in" if line.variance > 0 else "adjustment_out", quantity=line.variance, reference_type="stock_count", reference_id=count.id, reason=line.reason or "Approved physical count variance", performed_by=uid))
        count.status, count.approved_by, count.approved_at = "approved", uid, utcnow(); _audit(db, ctx, "approve", count)
        return as_dict(count)


def issue_job_part(ctx: dict, job_id: int, payload: dict) -> dict:
    wid, uid = ctx["workshop"]["id"], ctx["user"]["id"]
    item_id, quantity = int(payload.get("item_id") or 0), _number(payload.get("quantity"))
    if quantity <= 0:
        raise HTTPException(400, "Quantity must be greater than zero.")
    with SessionLocal.begin() as db:
        job = db.scalar(select(JobCard).where(JobCard.id == job_id, JobCard.workshop_id == wid))
        item = db.scalar(select(InventoryItem).where(InventoryItem.id == item_id, InventoryItem.workshop_id == wid))
        if not job or not item:
            raise HTTPException(404, "Job card or item not found.")
        balance = _balance(db, wid, job.branch_id, item_id)
        if balance.quantity_on_hand < quantity:
            raise HTTPException(409, "Insufficient stock.")
        line = db.scalar(select(JobCardItem).where(JobCardItem.job_card_id == job.id, JobCardItem.item_id == item_id))
        if not line:
            line = JobCardItem(workshop_id=wid, job_card_id=job.id, item_id=item_id, unit_price=item.selling_price)
            db.add(line); db.flush()
        released = min(balance.quantity_reserved, line.quantity_reserved, quantity)
        balance.quantity_reserved -= released
        line.quantity_reserved -= released
        balance.quantity_on_hand -= quantity
        line.quantity_issued += quantity
        movement = StockMovement(workshop_id=wid, branch_id=job.branch_id, item_id=item_id, movement_type="issue", quantity=-quantity, unit_cost=item.purchase_price, reference_type="job_card", reference_id=job.id, reason="Issued to repair job", performed_by=uid)
        db.add(movement); db.flush(); _audit(db, ctx, "issue_part", movement)
        return {"job_item": as_dict(line), "balance": as_dict(balance)}


def create_invoice(ctx: dict, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    lines = payload.get("lines") or []
    if not lines:
        raise HTTPException(400, "At least one invoice line is required.")
    with SessionLocal.begin() as db:
        branch_id, customer_id = int(payload.get("branch_id") or 0), int(payload.get("customer_id") or 0)
        branch = db.scalar(select(Branch).where(Branch.id == branch_id, Branch.workshop_id == wid))
        customer = db.scalar(select(Customer).where(Customer.id == customer_id, Customer.workshop_id == wid))
        vehicle_id = int(payload.get("vehicle_id") or 0)
        vehicle = db.scalar(select(Vehicle).where(Vehicle.id == vehicle_id, Vehicle.workshop_id == wid, Vehicle.customer_id == customer_id)) if vehicle_id else None
        job_id = int(payload.get("job_card_id") or 0)
        job = db.scalar(select(JobCard).where(JobCard.id == job_id, JobCard.workshop_id == wid)) if job_id else None
        if not branch or not customer or (vehicle_id and not vehicle) or (job_id and not job):
            raise HTTPException(404, "A selected branch, customer, vehicle, or job card is not valid for this workshop.")
        subtotal = ZERO; tax_total = ZERO
        invoice = Invoice(
            workshop_id=wid, branch_id=branch_id, customer_id=customer_id,
            vehicle_id=payload.get("vehicle_id") or None, job_card_id=payload.get("job_card_id") or None,
            invoice_number=str(payload.get("invoice_number") or f"INV-{utcnow():%Y%m%d%H%M%S}"),
            status=str(payload.get("status") or "issued"), due_on=date.fromisoformat(payload["due_on"]) if payload.get("due_on") else None,
        )
        db.add(invoice); db.flush()
        for raw in lines:
            quantity, unit_price, tax_rate = _number(raw.get("quantity", 1)), _number(raw.get("unit_price")), _number(raw.get("tax_rate"))
            base = quantity * unit_price; tax = base * tax_rate / 100; subtotal += base; tax_total += tax
            db.add(InvoiceLine(workshop_id=wid, invoice_id=invoice.id, line_type=str(raw.get("line_type") or "part"), item_id=raw.get("item_id") or None, description=str(raw.get("description") or "Workshop service"), quantity=quantity, unit_price=unit_price, tax_rate=tax_rate, line_total=base + tax))
        discount = _number(payload.get("discount_total"))
        invoice.subtotal, invoice.tax_total, invoice.discount_total = subtotal, tax_total, discount
        invoice.total = subtotal + tax_total - discount; invoice.balance_due = invoice.total
        _audit(db, ctx, "create", invoice)
        return as_dict(invoice)


def record_payment(ctx: dict, invoice_id: int, payload: dict) -> dict:
    wid = ctx["workshop"]["id"]
    amount = _number(payload.get("amount"))
    if amount <= 0:
        raise HTTPException(400, "Payment must be greater than zero.")
    with SessionLocal.begin() as db:
        invoice = db.scalar(select(Invoice).where(Invoice.id == invoice_id, Invoice.workshop_id == wid).with_for_update())
        if not invoice:
            raise HTTPException(404, "Invoice not found.")
        if amount > invoice.balance_due:
            raise HTTPException(409, "Payment exceeds the outstanding balance.")
        payment = Payment(workshop_id=wid, invoice_id=invoice.id, amount=amount, method=str(payload.get("method") or "cash"), reference=str(payload.get("reference") or ""))
        invoice.amount_paid += amount; invoice.balance_due -= amount
        if invoice.balance_due == 0:
            invoice.status = "paid"
        db.add(payment); db.flush(); _audit(db, ctx, "payment", payment)
        return {"payment": as_dict(payment), "invoice": as_dict(invoice)}


def dashboard(ctx: dict, branch_id: int | None = None) -> dict:
    wid = ctx["workshop"]["id"]
    with SessionLocal() as db:
        balance_filter = [StockBalance.workshop_id == wid]
        movement_filter = [StockMovement.workshop_id == wid]
        invoice_filter = [Invoice.workshop_id == wid]
        job_filter = [JobCard.workshop_id == wid]
        if branch_id:
            balance_filter.append(StockBalance.branch_id == branch_id); movement_filter.append(StockMovement.branch_id == branch_id)
            invoice_filter.append(Invoice.branch_id == branch_id); job_filter.append(JobCard.branch_id == branch_id)
        items = {x.id: x for x in db.scalars(select(InventoryItem).where(InventoryItem.workshop_id == wid)).all()}
        balances = db.scalars(select(StockBalance).where(*balance_filter)).all()
        inventory_value = sum((b.quantity_on_hand * items[b.item_id].purchase_price for b in balances if b.item_id in items), ZERO)
        retail_value = sum((b.quantity_on_hand * items[b.item_id].selling_price for b in balances if b.item_id in items), ZERO)
        low_stock = []
        for balance in balances:
            item = items.get(balance.item_id)
            if item and balance.quantity_on_hand - balance.quantity_reserved <= item.min_stock:
                low_stock.append({"item_id": item.id, "sku": item.sku, "name": item.name, "on_hand": float(balance.quantity_on_hand), "reserved": float(balance.quantity_reserved), "available": float(balance.quantity_on_hand - balance.quantity_reserved), "minimum": float(item.min_stock), "suggested_order": float(max(item.reorder_quantity, item.min_stock * 2 - (balance.quantity_on_hand - balance.quantity_reserved)))})
        since = utcnow() - timedelta(days=30)
        sales = db.scalar(select(func.coalesce(func.sum(Invoice.total), 0)).where(*invoice_filter, Invoice.issued_on >= since.date())) or 0
        receivables = db.scalar(select(func.coalesce(func.sum(Invoice.balance_due), 0)).where(*invoice_filter)) or 0
        open_jobs = db.scalar(select(func.count(JobCard.id)).where(*job_filter, JobCard.status.in_(["scheduled", "open", "in_progress"]))) or 0
        category_values = defaultdict(Decimal)
        for balance in balances:
            if balance.item_id in items:
                category_values[items[balance.item_id].category] += balance.quantity_on_hand * items[balance.item_id].purchase_price
        monthly = []
        for offset in range(5, -1, -1):
            current_month = date.today().year * 12 + date.today().month - 1 - offset
            month_start = date(current_month // 12, current_month % 12 + 1, 1)
            next_month = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
            value = db.scalar(select(func.coalesce(func.sum(Invoice.total), 0)).where(*invoice_filter, Invoice.issued_on >= month_start, Invoice.issued_on < next_month)) or 0
            monthly.append({"month": month_start.strftime("%b %Y"), "sales": float(value)})
        expiring = db.execute(
            select(StockBatch, InventoryItem.name).join(InventoryItem, InventoryItem.id == StockBatch.item_id)
            .where(StockBatch.workshop_id == wid, StockBatch.expires_on.is_not(None), StockBatch.expires_on <= date.today() + timedelta(days=60), StockBatch.quantity > 0)
            .order_by(StockBatch.expires_on).limit(20)
        ).all()
        top_usage = db.execute(
            select(InventoryItem.name, func.sum(-StockMovement.quantity).label("used"))
            .join(InventoryItem, InventoryItem.id == StockMovement.item_id)
            .where(*movement_filter, StockMovement.movement_type == "issue", StockMovement.occurred_at >= since)
            .group_by(InventoryItem.name).order_by(func.sum(-StockMovement.quantity).desc()).limit(8)
        ).all()
        return {
            "kpis": {"inventory_value": float(inventory_value), "retail_value": float(retail_value), "sales_30d": float(sales), "receivables": float(receivables), "low_stock_count": len(low_stock), "open_jobs": int(open_jobs)},
            "category_values": [{"category": k, "value": float(v)} for k, v in sorted(category_values.items(), key=lambda x: x[1], reverse=True)],
            "monthly_sales": monthly,
            "top_usage": [{"name": name, "quantity": float(value or 0)} for name, value in top_usage],
            "low_stock": sorted(low_stock, key=lambda x: x["available"] - x["minimum"]),
            "expiring": [{**as_dict(batch), "item_name": name} for batch, name in expiring],
        }


def recommendations(ctx: dict, branch_id: int | None = None) -> dict:
    stats = dashboard(ctx, branch_id)
    actions = []
    for item in stats["low_stock"][:10]:
        severity = "critical" if item["available"] <= 0 else "high"
        actions.append({"area": "inventory", "severity": severity, "title": f"Reorder {item['name']}", "recommendation": f"Order {item['suggested_order']:g} units. Available stock is {item['available']:g} against a minimum of {item['minimum']:g}.", "evidence": item, "requires_approval": True})
    for batch in stats["expiring"][:5]:
        actions.append({"area": "inventory", "severity": "high", "title": f"Use or return expiring {batch['item_name']}", "recommendation": f"Batch {batch['batch_number']} expires on {batch['expires_on']}. Prioritize it using FEFO or negotiate a supplier return.", "evidence": batch, "requires_approval": True})
    if stats["kpis"]["receivables"] > stats["kpis"]["sales_30d"] * 0.35 and stats["kpis"]["receivables"] > 0:
        actions.append({"area": "cashflow", "severity": "medium", "title": "Collect outstanding invoices", "recommendation": "Receivables are high relative to the last 30 days of sales. Review overdue invoices and send payment reminders.", "evidence": {"receivables": stats["kpis"]["receivables"], "sales_30d": stats["kpis"]["sales_30d"]}, "requires_approval": True})
    if stats["kpis"]["open_jobs"] > 10:
        actions.append({"area": "operations", "severity": "medium", "title": "Review workshop capacity", "recommendation": "The open job-card queue is elevated. Rebalance technician assignments and confirm promised delivery times.", "evidence": {"open_jobs": stats["kpis"]["open_jobs"]}, "requires_approval": True})
    if not actions:
        actions.append({"area": "operations", "severity": "info", "title": "No urgent exceptions", "recommendation": "Stock, receivables, and workshop workload show no rule-based critical exception. Continue daily movement capture for stronger forecasting.", "evidence": stats["kpis"], "requires_approval": False})
    result = {"generated_at": utcnow().isoformat(), "mode": "explainable_rules", "actions": actions, "disclaimer": "Recommendations are decision support. Purchases, transfers, price changes, and customer communications require a human approval."}
    model_id = os.getenv("INVENTORY_BEDROCK_MODEL_ID", "").strip()
    if model_id:
        try:
            import boto3
            prompt = "Summarize these automobile workshop KPIs and recommendations in under 180 words. Do not invent facts:\n" + json.dumps({"kpis": stats["kpis"], "actions": actions}, default=str)
            response = boto3.client("bedrock-runtime", region_name=os.getenv("AWS_REGION", "ap-south-1")).converse(modelId=model_id, messages=[{"role": "user", "content": [{"text": prompt}]}], inferenceConfig={"maxTokens": 300, "temperature": 0.1})
            result["executive_summary"] = response["output"]["message"]["content"][0]["text"]
            result["mode"] = "rules_plus_bedrock"
        except Exception as exc:
            result["ai_warning"] = f"Bedrock summary unavailable; explainable recommendations remain active: {exc}"
    return result


async def import_inventory(ctx: dict, branch_id: int, upload: UploadFile, dry_run: bool = True) -> dict:
    if not upload.filename or not upload.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Upload an .xlsx workbook.")
    data = await upload.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "Workbook exceeds the 10 MB limit.")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").strip().lower().replace(" ", "_") for cell in next(sheet.iter_rows())]
    except Exception as exc:
        raise HTTPException(400, "The workbook could not be read.") from exc
    required = {"sku", "name", "quantity"}
    if not required.issubset(headers):
        raise HTTPException(400, f"Required columns: {', '.join(sorted(required))}")
    rows, errors = [], []
    for index, values in enumerate(sheet.iter_rows(values_only=True), start=2):
        row = dict(zip(headers, values))
        if not any(value not in (None, "") for value in values):
            continue
        try:
            if not row.get("sku") or not row.get("name"):
                raise ValueError("SKU and name are required")
            row["quantity"] = float(row.get("quantity") or 0)
            rows.append(row)
        except Exception as exc:
            errors.append({"row": index, "error": str(exc)})
    if dry_run or errors:
        return {"status": "validated" if not errors else "invalid", "dry_run": True, "valid_rows": len(rows), "errors": errors, "preview": rows[:10], "verification_method": "Two-pass migration: validate first, then blind physical count and manager approval."}
    wid = ctx["workshop"]["id"]
    imported = 0
    for row in rows:
        with SessionLocal.begin() as db:
            item = db.scalar(select(InventoryItem).where(InventoryItem.workshop_id == wid, InventoryItem.sku == str(row["sku"])))
            if not item:
                item = InventoryItem(
                    workshop_id=wid, sku=str(row["sku"]), barcode=str(row.get("barcode") or row["sku"]), name=str(row["name"]),
                    category=str(row.get("category") or "General"), brand=str(row.get("brand") or ""), unit=str(row.get("unit") or "each"),
                    purchase_price=_number(row.get("purchase_price")), selling_price=_number(row.get("selling_price")), tax_rate=_number(row.get("tax_rate")),
                    min_stock=_number(row.get("min_stock")), reorder_quantity=_number(row.get("reorder_quantity")),
                )
                db.add(item); db.flush(); _audit(db, ctx, "import_create", item)
            balance = _balance(db, wid, branch_id, item.id)
            quantity = _number(row["quantity"])
            movement = StockMovement(workshop_id=wid, branch_id=branch_id, item_id=item.id, movement_type="opening", quantity=quantity, unit_cost=item.purchase_price, reference_type="excel_import", reason="Opening stock pending physical verification", performed_by=ctx["user"]["id"])
            balance.quantity_on_hand += quantity; db.add(movement); imported += 1
    return {"status": "imported", "rows": imported, "next_step": "Create a blind physical count and approve variances before treating opening stock as verified."}


def inventory_template() -> bytes:
    from openpyxl import Workbook
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Inventory"
    headers = ["sku", "barcode", "name", "category", "brand", "unit", "quantity", "purchase_price", "selling_price", "tax_rate", "min_stock", "reorder_quantity", "rack", "shelf", "bin_code", "batch_number", "expires_on", "warranty_until"]
    sheet.append(headers); sheet.append(["BP-001", "890000000001", "Brake Pad Front", "Brakes", "Example", "set", 10, 800, 1200, 18, 4, 8, "A", "01", "B02", "LOT-001", "2028-12-31", "2027-12-31"])
    for cell in sheet[1]: cell.font = cell.font.copy(bold=True)
    output = BytesIO(); workbook.save(output); return output.getvalue()


def barcode_labels(ctx: dict, item_ids: list[int]) -> bytes:
    wid = ctx["workshop"]["id"]
    with SessionLocal() as db:
        query = select(InventoryItem).where(InventoryItem.workshop_id == wid)
        if item_ids:
            query = query.where(InventoryItem.id.in_(item_ids))
        items = db.scalars(query.order_by(InventoryItem.name).limit(100)).all()
        if not items:
            raise HTTPException(404, "No items found for label printing.")
        output = BytesIO(); pdf = canvas.Canvas(output, pagesize=A4); width, height = A4
        label_w, label_h = width / 2, 92
        for index, item in enumerate(items):
            slot = index % 12; col, row = slot % 2, slot // 2
            if slot == 0 and index:
                pdf.showPage()
            x, y = col * label_w + 18, height - (row + 1) * label_h + 14
            pdf.setFont("Helvetica-Bold", 9); pdf.drawString(x, y + 55, item.name[:42])
            pdf.setFont("Helvetica", 7); pdf.drawString(x, y + 43, f"SKU: {item.sku}  |  {item.brand}")
            value = item.barcode or item.sku
            barcode = code128.Code128(value, barHeight=25, barWidth=0.75); barcode.drawOn(pdf, x, y + 10)
            pdf.setFont("Helvetica", 7); pdf.drawCentredString(x + 80, y + 2, value)
        pdf.save(); return output.getvalue()


def accounting_export(ctx: dict) -> bytes:
    """Low-cost, vendor-neutral export accepted as an import source by Tally/Zoho workflows."""
    wid = ctx["workshop"]["id"]
    output = StringIO(); writer = csv.writer(output)
    writer.writerow(["voucher_type", "voucher_number", "date", "party_id", "gross_amount", "tax_amount", "paid_amount", "balance_due", "status"])
    with SessionLocal() as db:
        invoices = db.scalars(select(Invoice).where(Invoice.workshop_id == wid).order_by(Invoice.issued_on, Invoice.id)).all()
        for invoice in invoices:
            writer.writerow(["Sales", invoice.invoice_number, invoice.issued_on.isoformat(), invoice.customer_id, invoice.total, invoice.tax_total, invoice.amount_paid, invoice.balance_due, invoice.status])
    return output.getvalue().encode("utf-8-sig")
